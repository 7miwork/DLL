#!/usr/bin/env python3
"""Test export functionality without running the full GUI."""

import sys
import os
import tempfile
from pathlib import Path

# Add the project root independently of the developer workstation.
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from PySide6.QtWidgets import QApplication, QTableWidgetItem
from main_window import MainWindow
import i18n


def test_csv_export():
    """Test CSV export directly without file dialog."""
    print("=== Test 1: Empty table ===")
    app = QApplication.instance() or QApplication(sys.argv)
    i18n.set_language("en")
    window = MainWindow()
    
    # Mock empty scans
    window._scans = []
    window._table.setRowCount(0)
    
    # Should show "no data" message
    window._export_scans()
    print("OK - empty table handled")
    
    # Test 2: CSV with special chars
    print("\n=== Test 2: CSV with special chars ===")
    window._scans.clear()
    window._table.setRowCount(0)
    
    test_data = [
        ("2026-08-04 14:30", "OP001", "GOOD", "ABC-123/日本語/中文", "", "saved"),
        ("2026-08-04 14:31", "OP002", "NG", "DEF-456", "", "failed"),
        ("2026-08-04 14:32", "OP001", "GOOD", "GHI-789", "", "saved"),
    ]
    
    for time_str, op, status, barcode, img, save_status in test_data:
        row = window._table.rowCount()
        window._table.insertRow(row)
        window._table.setItem(row, 0, QTableWidgetItem(time_str))
        window._table.setItem(row, 1, QTableWidgetItem(op))
        window._table.setItem(row, 2, QTableWidgetItem(status))
        window._table.setItem(row, 3, QTableWidgetItem(barcode))
        window._table.setItem(row, 4, QTableWidgetItem(img))
        window._table.setItem(row, 5, QTableWidgetItem(save_status))
    
    with tempfile.TemporaryDirectory() as tmpdir:
        csv_path = os.path.join(tmpdir, "test_export.csv")
        headers = [window._table.horizontalHeaderItem(col).text() for col in range(window._table.columnCount())]
        rows = []
        for row in range(window._table.rowCount()):
            row_data = [window._table.item(row, col).text() if window._table.item(row, col) else "" for col in range(window._table.columnCount())]
            rows.append(row_data)
        
        window._export_csv(csv_path, headers, rows)
        
        if os.path.exists(csv_path):
            # Read raw bytes to check BOM
            with open(csv_path, 'rb') as f:
                raw = f.read()
            has_bom = raw.startswith(b'\xef\xbb\xbf')
            print(f"CSV raw start: {raw[:20]!r}")
            text = raw.decode('utf-8-sig')
            print(f"CSV content:\n{text}")
            
            if has_bom:
                print("✓ UTF-8 BOM present")
            else:
                print("✗ UTF-8 BOM MISSING")
            
            if '日本語' in text and '中文' in text:
                print("✓ Special chars preserved")
            else:
                print("✗ Special chars MISSING")
            
            lines = text.strip().split("\n")
            print(f"✓ Total lines: {len(lines)} (1 header + {len(lines)-1} data)")
        else:
            print(f"✗ CSV not created at {csv_path}")
    
    # Test 3: XLSX with special chars
    print("\n=== Test 3: XLSX with special chars ===")
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = os.path.join(tmpdir, "test_export.xlsx")
        headers = [window._table.horizontalHeaderItem(col).text() for col in range(window._table.columnCount())]
        rows = []
        for row in range(window._table.rowCount()):
            row_data = [window._table.item(row, col).text() if window._table.item(row, col) else "" for col in range(window._table.columnCount())]
            rows.append(row_data)
        
        window._export_xlsx(xlsx_path, headers, rows)
        
        if os.path.exists(xlsx_path):
            from openpyxl import load_workbook
            wb = load_workbook(xlsx_path)
            ws = wb.active
            print(f"✓ XLSX created: {ws.title}, rows={ws.max_row}, cols={ws.max_column}")
            
            headers_out = [cell.value for cell in ws[1]]
            print(f"✓ Headers: {headers_out}")
            
            special_found = False
            for row in ws.iter_rows(values_only=True):
                if '日本語' in str(row) or '中文' in str(row):
                    special_found = True
                    break
            print("✓ Special chars preserved in XLSX" if special_found else "✗ Special chars MISSING in XLSX")
        else:
            print(f"✗ XLSX not created at {xlsx_path}")
    
    print("\n=== All export tests completed ===")
    window.close()
    app.processEvents()


if __name__ == "__main__":
    test_csv_export()
